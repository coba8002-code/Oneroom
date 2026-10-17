"""세무 준비 — 귀속연도 자료(판정 없음)와 지출 증빙."""
import io
import uuid
from datetime import date

import pytest
from openpyxl import load_workbook

from app.services import backup, tax_prep
from tests.conftest import SAMPLES

PNG = b"\x89PNG\r\n\x1a\n" + b"receipt" * 30
PDF = b"%PDF-1.4 tax invoice"


@pytest.fixture
def org(client, onboarded):
    o = onboarded["organization_id"]
    p = {"organization_id": o}
    client.post("/api/v1/charges/generate", params={**p, "period": "2026-11"})
    client.post(f"/api/v1/bank-accounts/{onboarded['bank_account_id']}/import", params=p,
                files={"file": ("s.csv", (SAMPLES / "bank_transactions_sample.csv").read_bytes())})
    client.post("/api/v1/matching/run", params=p)
    client.post("/api/v1/expenses/classify", params=p)
    return o


def _manual(client, org, **kw):
    body = {"category": "repair", "amount": 55000, "spent_on": "2026-11-10", "memo": "수전 교체", **kw}
    res = client.post("/api/v1/expenses", params={"organization_id": org}, json=body)
    assert res.status_code == 201, res.text
    return res.json()["id"]


def test_annual_marks_months_without_bank_data_and_does_not_fill(client, org):
    _manual(client, org)
    d = client.get("/api/v1/tax-prep/2026", params={"as_of": "2026-12-31"}).json()
    assert "신고 확정이 아니며" in d["notice"]
    st = d["status"]
    assert st["months_with_bank_data"] == ["2026-11"]  # 샘플 은행 내역은 11월만
    assert "2026-01" in st["months_without_bank_data"]
    jan = d["months"][0]
    assert jan["bank_covered"] is False and jan["rent_received"] is None and jan["operating_total"] is None
    nov = next(m for m in d["months"] if m["period"] == "2026-11")
    assert nov["rent_charged"] > 0 and nov["rent_received"] is not None
    t = d["totals_known_months"]
    assert t["rent_received"] == nov["rent_received"]  # 자료 있는 달만 더함
    assert t["by_category"]["repair"] >= 55000
    assert st["operating_without_receipt"] >= 1
    assert d["deposits_total"] > 0 and all(x["returned_at"] is None for x in d["deposits"])
    assert client.get("/api/v1/tax-prep/1999", params={"as_of": "2026-12-31"}).status_code == 409


def test_receipts_attach_view_block_delete_and_backup(client, org, session, make_client):
    eid = _manual(client, org)
    url = f"/api/v1/expenses/{eid}/receipts"
    r = client.post(url, files={"file": ("영수증.png", PNG)})
    assert r.status_code == 201
    assert client.post(url, files={"file": ("again.png", PNG)}).status_code == 409
    assert client.post(url, files={"file": ("tax.pdf", PDF)}).status_code == 201
    assert client.post(url, files={"file": ("x.txt", b"hello")}).status_code == 415
    rid = r.json()["id"]
    got = client.get(f"/api/v1/expense-receipts/{rid}")
    assert got.status_code == 200 and got.content == PNG
    other, _ = make_client(org_name="남")
    assert other.get(f"/api/v1/expense-receipts/{rid}").status_code == 404
    exp = next(e for e in client.get("/api/v1/tax-prep/2026", params={"as_of": "2026-12-31"}).json()["expenses"] if e["id"] == eid)
    assert [x["name"] for x in exp["receipts"]] == ["영수증.png", "tax.pdf"]
    # 증빙이 붙은 지출은 지우지 못함
    res = client.delete(f"/api/v1/expenses/{eid}", params={"organization_id": org})
    assert res.status_code == 409 and "증빙" in res.json()["detail"]
    names = backup.file_names(backup.collect(session, uuid.UUID(org)))
    assert len([n for n in names]) >= 2
    for x in exp["receipts"]:
        assert client.delete(f"/api/v1/expense-receipts/{x['id']}").status_code == 204
    assert client.delete(f"/api/v1/expenses/{eid}", params={"organization_id": org}).status_code == 204


def test_xlsx_export_has_sheets_and_status(client, org):
    _manual(client, org)
    res = client.get("/api/v1/tax-prep/2026/export.xlsx", params={"as_of": "2026-12-31"})
    assert res.status_code == 200 and "roomax-tax-2026.xlsx" in res.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["요약", "월별", "경비 내역", "보증금"]
    summary = {r[0].value: r[1].value for r in wb["요약"].iter_rows() if r[0].value}
    assert summary["귀속연도"] == 2026 and "신고 확정이 아니며" in summary["주의"]
    assert "2026-01" in summary["은행 내역 없는 달 (합계에서 빠짐)"]
    jan = [c.value for c in wb["월별"][2]]
    assert jan[:5] == ["2026-01", "없음", "—", "—", "—"]
    assert any(r[1].value == "수리비" and r[3].value == 55000 for r in wb["경비 내역"].iter_rows(min_row=2))


def test_annual_service_future_months_and_year_end_deposits(session, org):
    d = tax_prep.annual(session, uuid.UUID(org), 2026, date(2026, 11, 20))
    assert d["months"][11]["future"] is True and d["status"]["year_complete"] is False
    assert d["deposits_at"] == date(2026, 11, 20)
