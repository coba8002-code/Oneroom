"""월간 운영 보고서 — 장부 값 그대로, 모르면 비움, 이름 없음, 엑셀."""
import io
import re
from datetime import date

import pytest
from openpyxl import load_workbook

from app.routers import monthly_report as router_mod
from tests.conftest import SAMPLES


@pytest.fixture(autouse=True)
def _today(monkeypatch):
    monkeypatch.setattr(router_mod, "_today", lambda: date(2026, 12, 15))


@pytest.fixture
def org(client, onboarded):
    o = onboarded["organization_id"]
    client.post("/api/v1/charges/generate", params={"organization_id": o, "period": "2026-11"})
    raw = (SAMPLES / "bank_transactions_sample.csv").read_bytes()
    client.post(f"/api/v1/bank-accounts/{onboarded['bank_account_id']}/import", files={"file": ("s.csv", raw)})
    client.post("/api/v1/matching/run", params={"organization_id": o})
    return o


def _get(client, period):
    return client.get("/api/v1/reports/monthly-summary", params={"period": period})


def test_matches_ledger(client, org):
    r = _get(client, "2026-11").json()
    base = client.get("/api/v1/reports/monthly", params={"organization_id": org, "period": "2026-11"}).json()
    assert r["as_of"] == "2026-11-30" and r["month_closed"] is True
    assert r["money"]["rent_charged"] == base["rent_charged"] and r["money"]["noi"] == base["noi"]
    assert r["status"]["bank_covered"] is True and r["status"]["pending_review_in"] == 4
    assert {i["unit_no"]: i["outstanding"] for i in r["overdue"]["items"]} == {"103": 420000, "301": 480000, "302": 450000}
    assert r["overdue"]["total"] == 1_350_000
    assert len(r["charges"]) == 8 and r["occupancy"] == {"rental_units": 9, "leased_in_month": 8}
    assert r["deposits_liability"] == 43_000_000  # 샘플 보증금 합계 (모두 수령, 반환 없음)
    assert not re.search(r"임차인[A-H]|010-", str(r))  # 이름·연락처 없음


def test_month_without_bank_data_is_blank(client, org):
    r = _get(client, "2026-10").json()
    assert r["status"]["bank_covered"] is False
    assert r["money"]["rent_received"] is None and r["money"]["operating_total"] is None and r["money"]["noi"] is None
    assert r["money"]["rent_charged"] is None  # 그 달 청구를 만들지 않음 → 모름


def test_current_month_uses_today_and_future_rejected(client, org):
    assert _get(client, "2026-12").json()["as_of"] == "2026-12-15"
    assert _get(client, "2026-12").json()["month_closed"] is False
    assert _get(client, "2027-01").status_code == 422
    assert _get(client, "2026-13").status_code == 422


def test_xlsx(client, org):
    res = client.get("/api/v1/reports/monthly-summary/export.xlsx", params={"period": "2026-11"})
    assert res.status_code == 200 and "roomax-report-2026-11.xlsx" in res.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["요약", "청구", "미수", "공실", "점검·수리"]
    summary = {row[0]: row[1] for row in wb["요약"].iter_rows(values_only=True) if row and row[0]}
    assert summary["미수 (기준일, 납부일 지남)"] == 1_350_000 and summary["보증금 잔액 (부채, 월말)"] == 43_000_000
    oct_wb = load_workbook(io.BytesIO(client.get("/api/v1/reports/monthly-summary/export.xlsx",
                                                 params={"period": "2026-10"}).content))
    oct_summary = {row[0]: row[1] for row in oct_wb["요약"].iter_rows(values_only=True) if row and row[0]}
    assert oct_summary["순운영수익"] == "—"  # 모르면 빈칸 표시
