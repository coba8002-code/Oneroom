"""월간 운영 보고서 — 한 달을 한 장으로 (화면 인쇄·PDF 저장, 엑셀).

모든 숫자는 장부 조회 결과 (CLAUDE.md 규칙 7). 모르는 값은 None → 화면 '—'.
- 결산(청구·입금·운영비·순운영수익)은 reports.monthly 그대로 — 은행 내역이 그 달과 겹치지 않으면 비움
- 미수는 '기준일'(그 달 말일, 진행 중인 달이면 오늘) 기준 납부일 지난 것. 검토 대기·은행 내역 기간을 함께 적는다
- 그 달 계약이 있던 호실 수는 계약 기간(시작일~실제 종료일)으로 센다 — 현재 상태로 과거 달을 짐작하지 않음
- 임차인 이름·연락처는 넣지 않는다 (보고서를 세무사 등에게 넘길 수 있으므로 호실로만)
"""
from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import and_, or_, select
from sqlalchemy.orm import Session

from app.models import Charge, Deposit, Lease, Property, RoutineLog, RoutineTask, Unit
from app.services import dashboard, fixtures, ledger, receivables, reports, routines, vacancy
from app.services.dashboard import LABEL as LABELS
from app.services.expenses import NON_OPERATING, OPERATING

NOTICE = "장부에 기록된 값만 모았습니다. 비어 있는 칸(—)은 자료가 없어 계산하지 않은 값입니다. 세금 신고용 확정 자료가 아닙니다."


class ReportError(ValueError):
    pass


def _org_units(organization_id: uuid.UUID):
    return (select(Unit).join(Property, Unit.property_id == Property.id)
            .where(Property.organization_id == organization_id))


def build(session: Session, organization_id: uuid.UUID, period: str, today: date) -> dict:
    try:
        first, nxt, _, _ = reports.month_range(period)
    except ValueError as e:
        raise ReportError("월은 YYYY-MM 형식입니다") from e
    if first > today:
        raise ReportError("아직 오지 않은 달입니다")
    month_end = nxt - timedelta(days=1)
    as_of = min(month_end, today)

    m = reports.monthly(session, organization_id, period)
    recv = receivables.compute(session, organization_id, as_of)

    units = session.scalars(_org_units(organization_id).where(Unit.use_type == "rental")).all()
    leased = session.scalars(
        select(Lease.unit_id).where(
            Lease.unit_id.in_([u.id for u in units]), Lease.status != "draft", Lease.start_date <= month_end,
            or_(Lease.ended_on.is_(None), Lease.ended_on >= first))).all()

    charges = []
    for c, label in session.execute(
            select(Charge, Unit.label).join(Lease, Charge.lease_id == Lease.id).join(Unit, Lease.unit_id == Unit.id)
            .join(Property, Unit.property_id == Property.id)
            .where(Property.organization_id == organization_id, Charge.period == period)
            .order_by(Unit.label)).all():
        paid = ledger.paid_amount(session, c.id)
        charges.append({"unit_no": label, "due_date": c.due_date, "amount": c.amount, "paid": int(paid),
                        "status": c.status, "kind": c.kind})

    unreturned = session.scalars(
        select(Deposit.amount).join(Lease, Deposit.lease_id == Lease.id).join(Unit, Lease.unit_id == Unit.id)
        .join(Property, Unit.property_id == Property.id)
        .where(Property.organization_id == organization_id, Deposit.received_at.is_not(None),
               Deposit.received_at <= month_end,
               or_(Deposit.returned_at.is_(None), Deposit.returned_at > month_end))).all()

    done = session.execute(
        select(RoutineLog, RoutineTask.title, Unit.label)
        .join(RoutineTask, RoutineLog.task_id == RoutineTask.id)
        .join(Property, RoutineTask.property_id == Property.id)
        .outerjoin(Unit, RoutineTask.unit_id == Unit.id)
        .where(Property.organization_id == organization_id,
               and_(RoutineLog.done_on >= first, RoutineLog.done_on <= month_end))
        .order_by(RoutineLog.done_on)).all()
    overdue_tasks = [t for t in routines.list_tasks(session, organization_id, as_of) if t["state"] == "overdue"]

    repairs = [r for r in fixtures.list_repairs(session, organization_id)
               if first <= r["reported_on"] <= month_end or (r["completed_on"] and first <= r["completed_on"] <= month_end)]

    vac = vacancy.summary(session, organization_id, as_of)
    endings = [{k: v for k, v in e.items() if k != "tenant_name"}
               for e in dashboard.build(session, organization_id, as_of)["lease_endings"]]

    overdue = [i for i in recv.items if i.days_past_due > 0]
    return {
        "period": period, "as_of": as_of, "month_closed": month_end < today, "generated_at": datetime.now(timezone.utc),
        "notice": NOTICE,
        "status": {"bank_covered": m["bank_covered"], "bank_data_as_of": m["bank_data_as_of"],
                   "charges_ready": m["charges_ready"], "unconfirmed_in": m["unconfirmed_in"],
                   "unconfirmed_out": m["unconfirmed_out"], "pending_review_in": recv.pending_review,
                   "complete": m["complete"]},
        "occupancy": {"rental_units": len(units), "leased_in_month": len(set(leased))},
        "money": {"rent_charged": m["rent_charged"], "rent_received": m["rent_received"],
                  "mgmt_charged": m["mgmt_charged"], "mgmt_received": m["mgmt_received"],
                  "utility_charged": m["utility_charged"], "utility_received": m["utility_received"],
                  "operating": {LABELS.get(k, k): v for k, v in m["operating"].items()},
                  "operating_total": m["operating_total"], "noi": m["noi"],
                  "excluded": {LABELS.get(k, k): m["excluded"][k] for k in NON_OPERATING},
                  "deposit_in": m["excluded"]["deposit_in"], "deposit_out": m["excluded"]["deposit_out"]},
        "charges": charges,
        "overdue": {"total": int(sum(i.outstanding for i in overdue)),
                    "items": [{"unit_no": i.unit_no, "period": i.period, "kind": i.kind, "due_date": i.due_date,
                               "outstanding": i.outstanding, "days_past_due": i.days_past_due} for i in overdue]},
        "deposits_liability": int(sum(unreturned)),
        "vacancies": [{k: u[k] for k in ("unit_no", "vacant_since", "days_vacant", "missed_rent", "asking_rent")}
                      for u in vac["units"]],
        "routines_done": [{"done_on": g.done_on, "title": title, "unit_no": label, "cost": g.cost}
                          for g, title, label in done],
        "routines_overdue": [{"title": t["title"], "unit_no": t["unit_no"], "next_due_on": t["next_due_on"]}
                             for t in overdue_tasks],
        "repairs": [{k: r[k] for k in ("unit_no", "fixture_name", "reported_on", "completed_on", "status", "cost", "payer")}
                    for r in repairs],
        "lease_endings": endings,
        "operating_categories": [LABELS.get(c, c) for c in OPERATING],
    }


def _v(x):
    return "—" if x is None else x


def to_xlsx(r: dict) -> bytes:
    wb = Workbook()
    bold = Font(bold=True)
    ws = wb.active
    ws.title = "요약"
    st, mo = r["status"], r["money"]
    rows = [
        ["ROOMAX 월간 운영 보고서", r["period"]], ["기준일", r["as_of"].isoformat()],
        ["작성 시각 (KST)", r["generated_at"].astimezone(reports.KST).strftime("%Y-%m-%d %H:%M")],
        ["주의", r["notice"]], [],
        ["자료 상태", ""],
        ["은행 내역 반영", "있음" if st["bank_covered"] else "없음 — 입금·운영비 비움"],
        ["은행 내역 최신 시각", st["bank_data_as_of"].astimezone(reports.KST).strftime("%Y-%m-%d %H:%M") if st["bank_data_as_of"] else "—"],
        ["미확정 거래 (입금/출금)", f'{st["unconfirmed_in"]} / {st["unconfirmed_out"]}'],
        ["입금 검토 대기 (전체)", st["pending_review_in"]], [],
        ["임대 호실 / 그 달 계약 있던 호실", f'{r["occupancy"]["rental_units"]} / {r["occupancy"]["leased_in_month"]}'],
        ["월세 청구", _v(mo["rent_charged"])], ["월세 입금 (확정)", _v(mo["rent_received"])],
        ["관리비 청구", _v(mo["mgmt_charged"])], ["관리비 입금 (확정)", _v(mo["mgmt_received"])],
        ["공과금 청구 (배분)", _v(mo["utility_charged"])], ["공과금 입금 (확정)", _v(mo["utility_received"])],
        ["미수 (기준일, 납부일 지남)", r["overdue"]["total"]],
        *[[f"운영비 — {k}", _v(v)] for k, v in mo["operating"].items()],
        ["운영비 합계", _v(mo["operating_total"])], ["순운영수익", _v(mo["noi"])],
        *[[f"제외 — {k}", _v(v)] for k, v in mo["excluded"].items()],
        ["보증금 받음 / 돌려줌 (그 달)", f'{_v(mo["deposit_in"])} / {_v(mo["deposit_out"])}'],
        ["보증금 잔액 (부채, 월말)", r["deposits_liability"]],
    ]
    for row in rows:
        ws.append(row)
    for cell in ws["A"]:
        cell.font = bold
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 40

    def sheet(title: str, header: list[str], data: list[list]):
        s = wb.create_sheet(title)
        s.append(header)
        for c in s[1]:
            c.font = bold
        for row in data:
            s.append([_v(x) for x in row])

    sheet("청구", ["호실", "납부일", "청구액", "입금", "상태"],
          [[c["unit_no"], c["due_date"].isoformat(), c["amount"], c["paid"], c["status"]] for c in r["charges"]])
    sheet("미수", ["호실", "청구 월", "납부일", "남은 금액", "지난 일수"],
          [[i["unit_no"], i["period"], i["due_date"].isoformat(), i["outstanding"], i["days_past_due"]] for i in r["overdue"]["items"]])
    sheet("공실", ["호실", "공실 시작", "공실 일수", "놓친 월세(직전 조건)", "모집 월세"],
          [[v["unit_no"], v["vacant_since"].isoformat() if v["vacant_since"] else None, v["days_vacant"], v["missed_rent"],
            v["asking_rent"]] for v in r["vacancies"]])
    sheet("점검·수리", ["구분", "날짜", "내용", "위치", "비용", "상태"],
          [["정기 점검", d["done_on"].isoformat(), d["title"], d["unit_no"] or "공용", d["cost"], "완료"] for d in r["routines_done"]]
          + [["정기 점검", t["next_due_on"].isoformat(), t["title"], t["unit_no"] or "공용", None, "예정일 지남"] for t in r["routines_overdue"]]
          + [["수리", x["reported_on"].isoformat(), x["fixture_name"] or "수리", x["unit_no"] or "공용", x["cost"],
              "완료" if x["status"] == "done" else "진행 중"] for x in r["repairs"]])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
