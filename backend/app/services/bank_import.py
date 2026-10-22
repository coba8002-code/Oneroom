"""은행 거래내역 CSV/XLSX 가져오기.

중복 방지 (PRD §8 "동일 파일 재업로드 시 거래 중복 없음")
- 은행 거래고유번호가 있으면 그것으로 fingerprint.
- 없으면 (일시·방향·금액·잔액·적요·상대방) + "같은 파일 안에서 같은 내용이 몇 번째인지".
  → 같은 날 같은 금액이 정말 2건이면 둘 다 들어가고, 같은 파일을 다시 올리면 둘 다 중복으로 걸러진다.
- (bank_account_id, fingerprint) 유일 제약이 최종 방어선.
- 휴대폰 입금 알림으로 먼저 들어온 거래(source='alert')는 같은 계좌·방향·금액이 30분 안(날짜만 있는 파일은
  같은 날)에 있으면 새로 넣지 않고 그 거래를 내역 값으로 고쳐 합친다 — 이미 한 확정·배분은 그대로 남는다.

은행별 열 이름은 실물 파일 확인 후 COLUMN_ALIASES 에 추가한다.
- 우리은행 '계좌거래내역 조회'(2026-10 실물 확인): 거래일시 · 적요(거래 방식) · 기재내용(보낸/받는 사람) ·
  지급금액 · 입금금액 · 잔액 · 취급점. 열 이름 끝의 '(원)'·공백은 무시한다.
- 파일: CSV · XLSX · XLS(옛 엑셀, 은행이 HTML 표를 .xls 로 주는 경우 포함).
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import uuid
from collections import Counter
from html.parser import HTMLParser
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from zoneinfo import ZoneInfo

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import AuditLog, BankAccount, Transaction
from app.schemas import ImportOut, RowError
from app.services import matching
from app.services.onboarding import decode_text

KST = ZoneInfo("Asia/Seoul")

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "occurred_at": ("거래일시", "거래일자", "거래일", "일시", "date"),
    "time": ("거래시간", "시간", "time"),
    "amount_in": ("입금", "입금액", "맡기신금액", "입금금액", "deposit"),
    "amount_out": ("출금", "출금액", "찾으신금액", "출금금액", "지급", "지급금액", "withdrawal"),
    "description": ("내용", "적요", "거래내용", "description"),
    # 우리은행은 '기재내용'에 입금자 이름이 들어간다 — 입금자명으로 계약을 찾으려면 상대방이어야 한다
    "counterparty": ("입금자", "보낸분/받는분", "의뢰인", "상대방", "기재내용", "counterparty"),
    "balance": ("잔액", "거래후잔액", "balance"),
    "bank_ref": ("거래고유번호", "거래번호", "bank_ref"),
}
REQUIRED = ("occurred_at", "amount_in", "amount_out")
DATETIME_FORMATS = (
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
    "%Y.%m.%d %H:%M:%S", "%Y.%m.%d %H:%M", "%Y.%m.%d",
    "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y/%m/%d",
    "%Y%m%d",
)


@dataclass
class ParsedRow:
    occurred_at: datetime  # UTC
    direction: str
    amount: int
    description: str | None
    counterparty: str | None
    balance_after: int | None
    bank_ref: str | None


class _HtmlTable(HTMLParser):
    """은행이 .xls 이름으로 주는 HTML 표 → 행 목록 (tr/td/th 만 본다)."""

    def __init__(self):
        super().__init__()
        self.rows: list[list[str]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def read_table(filename: str, raw: bytes) -> list[list]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith(".xls"):
        if raw[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":  # 진짜 옛 엑셀(BIFF)
            import xlrd

            book = xlrd.open_workbook(file_contents=raw)
            sh = book.sheet_by_index(0)
            return [[xlrd.xldate_as_datetime(c.value, book.datemode) if c.ctype == xlrd.XL_CELL_DATE else c.value
                     for c in sh.row(i)] for i in range(sh.nrows)]
        parser = _HtmlTable()  # 이름만 .xls 인 HTML 표
        parser.feed(decode_text(raw))
        return parser.rows
    return list(csv.reader(io.StringIO(decode_text(raw))))


def _norm_header(cell) -> str:
    """'입금금액(원)', '거래후 잔액' → '입금금액', '거래후잔액'"""
    return re.sub(r"\((원|KRW)\)$", "", re.sub(r"\s+", "", str(cell))) if cell is not None else ""


def _find_header(table: list[list]) -> tuple[int, dict[str, int]] | None:
    """은행 파일은 위쪽에 계좌정보 몇 줄이 붙는 경우가 많다 → 앞 20행에서 헤더를 찾는다."""
    for idx, row in enumerate(table[:20]):
        cells = [_norm_header(c) for c in row]
        mapping: dict[str, int] = {}
        for key, aliases in COLUMN_ALIASES.items():
            for i, cell in enumerate(cells):
                if cell in aliases:
                    mapping[key] = i
                    break
        if all(k in mapping for k in REQUIRED):
            return idx, mapping
    return None


def _to_int(v) -> int | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    s = str(v).replace(",", "").replace("원", "").strip()
    if not s or s == "-":
        return None
    return int(round(float(s)))


def _to_utc(v, t=None) -> datetime:
    if isinstance(v, datetime):
        local = v
    elif isinstance(v, date):
        local = datetime.combine(v, time())
    else:
        s = str(v).strip()
        for fmt in DATETIME_FORMATS:
            try:
                local = datetime.strptime(s, fmt)
                break
            except ValueError:
                continue
        else:
            raise ValueError(f"날짜를 읽을 수 없습니다: {s}")
    if t not in (None, ""):
        tt = t if isinstance(t, time) else time.fromisoformat(str(t).strip())
        local = datetime.combine(local.date(), tt)
    if local.tzinfo is None:
        local = local.replace(tzinfo=KST)  # 은행 내역은 한국시간
    return local.astimezone(timezone.utc)


def parse_rows(table: list[list]) -> tuple[list[tuple[int, ParsedRow]], list[RowError], int]:
    found = _find_header(table)
    if not found:
        names = ", ".join(f"{k}({'/'.join(v[:3])})" for k, v in COLUMN_ALIASES.items() if k in REQUIRED)
        return [], [RowError(row=1, message=f"헤더를 찾지 못했습니다. 필요한 열: {names}")], 0
    header_idx, col = found

    def cell(row: list, key: str):
        i = col.get(key)
        if i is None or i >= len(row):
            return None
        v = row[i]
        return v.strip() if isinstance(v, str) else v

    parsed: list[tuple[int, ParsedRow]] = []
    errors: list[RowError] = []
    total = 0
    for idx in range(header_idx + 1, len(table)):
        row = table[idx]
        if not any(c not in (None, "") for c in row):
            continue
        total += 1
        line = idx + 1
        try:
            a_in = _to_int(cell(row, "amount_in")) or 0
            a_out = _to_int(cell(row, "amount_out")) or 0
            if (a_in > 0) == (a_out > 0):
                raise ValueError("입금·출금 중 정확히 하나에만 금액이 있어야 합니다")
            ref = cell(row, "bank_ref")
            parsed.append((line, ParsedRow(
                occurred_at=_to_utc(cell(row, "occurred_at"), cell(row, "time")),
                direction="in" if a_in else "out",
                amount=a_in or a_out,
                description=(str(cell(row, "description")) if cell(row, "description") else None),
                counterparty=(str(cell(row, "counterparty")) if cell(row, "counterparty") else None),
                balance_after=_to_int(cell(row, "balance")),
                bank_ref=str(ref) if ref not in (None, "") else None,
            )))
        except (ValueError, TypeError) as e:
            errors.append(RowError(row=line, message=str(e)))
    return parsed, errors, total


def fingerprints(rows: list[ParsedRow]) -> list[str]:
    seen: Counter[str] = Counter()
    out: list[str] = []
    for r in rows:
        if r.bank_ref:
            base = f"ref|{r.bank_ref}"
        else:
            base = "|".join(str(x) for x in (
                r.occurred_at.isoformat(), r.direction, r.amount, r.balance_after,
                r.description or "", r.counterparty or "",
            ))
            seen[base] += 1
            base = f"{base}|#{seen[base]}"
        out.append(hashlib.sha256(base.encode()).hexdigest())
    return out


def import_transactions(session: Session, account: BankAccount, filename: str, raw: bytes) -> ImportOut:
    batch = uuid.uuid4().hex
    parsed, errors, total = parse_rows(read_table(filename, raw))
    fps = fingerprints([p for _, p in parsed])

    existing = set(session.scalars(
        select(Transaction.fingerprint).where(
            Transaction.bank_account_id == account.id, Transaction.fingerprint.in_(fps)
        )
    )) if fps else set()

    inserted = duplicates = merged = 0
    used: set = set()
    for (_, r), fp in zip(parsed, fps):
        if fp in existing:
            duplicates += 1
            continue
        existing.add(fp)  # 같은 파일 안 거래고유번호 중복도 방어
        if alert := _alert_twin(session, account.id, r, used):
            used.add(alert.id)
            alert.occurred_at, alert.description, alert.counterparty = r.occurred_at, r.description, r.counterparty
            alert.balance_after, alert.fingerprint, alert.import_batch = r.balance_after, fp, batch
            alert.source = "statement"
            merged += 1
            continue
        session.add(Transaction(
            bank_account_id=account.id, occurred_at=r.occurred_at, amount=r.amount,
            direction=r.direction, description=r.description, counterparty=r.counterparty,
            balance_after=r.balance_after, fingerprint=fp, import_batch=batch,
        ))
        inserted += 1

    session.flush()
    leftover = 0
    if parsed:
        first = min(r.occurred_at for _, r in parsed)
        last = max(r.occurred_at for _, r in parsed)
        last += timedelta(days=1) if last.astimezone(KST).time() == time() else ALERT_WINDOW  # 날짜만 있으면 그날 끝까지
        leftover = len(session.scalars(select(Transaction.id).where(
            Transaction.bank_account_id == account.id, Transaction.source == "alert",
            Transaction.occurred_at >= first, Transaction.occurred_at < last)).all())
    session.add(AuditLog(
        action="transactions.import", entity_type="bank_account", entity_id=str(account.id),
        detail={"batch": batch, "file_type": filename.rsplit(".", 1)[-1].lower(), "total": total,
                "inserted": inserted, "duplicates": duplicates, "errors": len(errors),
                "merged_alerts": merged, "alerts_without_statement": leftover},
    ))
    session.commit()
    return ImportOut(batch=batch, total_rows=total, inserted=inserted, duplicates=duplicates, errors=errors,
                     merged_alerts=merged, alerts_without_statement=leftover)


ALERT_WINDOW = timedelta(minutes=30)


def _alert_twin(session: Session, account_id: uuid.UUID, r: ParsedRow, used: set) -> Transaction | None:
    """이 내역 줄과 같은 거래로 보이는 입금 알림 거래. 후보가 여럿이면 입금자명이 같은 것 → 가까운 시각 순."""
    local = r.occurred_at.astimezone(KST)
    if local.time() == time():  # 시각 없이 날짜만 있는 파일 → 같은 날(KST)
        lo, hi = r.occurred_at, r.occurred_at + timedelta(days=1)
    else:
        lo, hi = r.occurred_at - ALERT_WINDOW, r.occurred_at + ALERT_WINDOW
    cands = [t for t in session.scalars(select(Transaction).where(
        Transaction.bank_account_id == account_id, Transaction.source == "alert",
        Transaction.direction == r.direction, Transaction.amount == r.amount,
        Transaction.occurred_at >= lo, Transaction.occurred_at < hi)) if t.id not in used]
    if not cands:
        return None
    name = matching.normalize_name(r.counterparty)
    when = lambda t: abs((t.occurred_at if t.occurred_at.tzinfo else t.occurred_at.replace(tzinfo=timezone.utc))  # noqa: E731
                         - r.occurred_at)
    return min(cands, key=lambda t: (matching.normalize_name(t.counterparty) != name, when(t)))
